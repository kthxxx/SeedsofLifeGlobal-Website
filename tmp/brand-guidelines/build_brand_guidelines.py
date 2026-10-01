from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Projects\seedsoflife-global")
REFERENCE = Path(r"C:\Users\keith\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-strategy-memorandum\assets\reference.docx")
OUTPUT = ROOT / "Seeds_of_Life_Global_Brand_Guidelines.docx"
FULL_LOGO = ROOT / "public" / "brand" / "seeds-of-life-global-full-color.png"
ICON_COLOR = ROOT / "public" / "brand" / "seeds-of-life-global-icon-color.png"
ICON_FOREST = ROOT / "public" / "brand" / "seeds-of-life-global-icon-forest.png"
ICON_WHITE = ROOT / "public" / "brand" / "seeds-of-life-global-icon-white.png"

FOREST = "467A63"
DEEP_FOREST = "173B2A"
SEED_GOLD = "FFDE59"
CREAM = "F7F3E9"
MOSS = "58724D"
CLAY = "A75434"
PALE_LEAF = "DCE5CF"
CHARCOAL = "243229"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=FOREST, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = OxmlElement("w:keepNext")
    p_pr.append(keep_next)


def set_spacing(paragraph, before=0, after=6, line=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    if line:
        paragraph.paragraph_format.line_spacing = line


def add_run(paragraph, text, *, bold=False, italic=False, size=None, color=None, font="Arial"):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_para(doc, text="", *, style=None, bold_prefix=None, color=None, size=None, align=None, before=0, after=6, line=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        add_run(p, bold_prefix, bold=True, color=color, size=size)
        add_run(p, text[len(bold_prefix):], color=color, size=size)
    else:
        add_run(p, text, color=color, size=size)
    set_spacing(p, before, after, line)
    return p


def add_heading(doc, text, level=1):
    sizes = {1: 26, 2: 17, 3: 12}
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"] if level == 1 else doc.styles["Heading 2"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_run(p, text, bold=True, size=sizes[level], color=DEEP_FOREST, font="Georgia")
    set_spacing(p, before=10 if level == 1 else 8, after=7)
    set_keep(p)
    return p


def add_kicker(doc, text):
    p = doc.add_paragraph()
    add_run(p, text.upper(), bold=True, size=8, color=FOREST)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    return p


def add_rule(doc, color=SEED_GOLD):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    cell.height = Inches(0.055)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22 + (0.18 * level))
    p.paragraph_format.first_line_indent = Inches(-0.16)
    add_run(p, "•  ", bold=True, size=10.3, color=FOREST)
    add_run(p, text, size=10.3, color=CHARCOAL)
    set_spacing(p, after=3, line=1.08)
    return p


def add_callout(doc, label, body, *, fill=PALE_LEAF, label_color=DEEP_FOREST):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=fill, size="4")
    set_cell_margins(cell, top=145, start=180, bottom=145, end=180)
    p = cell.paragraphs[0]
    add_run(p, label.upper() + "  ", bold=True, size=8.4, color=label_color)
    add_run(p, body, size=9.8, color=CHARCOAL)
    set_spacing(p, after=0, line=1.08)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_color_swatch_table(doc):
    rows = [
        ("Primary forest", FOREST, "#467A63", "Core brand green; primary field and key identifier."),
        ("Seed gold", SEED_GOLD, "#FFDE59", "Warm energy and growth; use as a bright accent."),
        ("Deep forest", DEEP_FOREST, "#173B2A", "High-contrast copy, navigation, and grounded depth."),
        ("Warm cream", CREAM, "#F7F3E9", "Primary light background; calm and inviting."),
        ("Moss", MOSS, "#58724D", "Secondary organic accent; sparing use."),
        ("Clay", CLAY, "#A75434", "Human warmth for calls to action and emphasis."),
        ("Pale leaf", PALE_LEAF, "#DCE5CF", "Soft panel fill and supporting background."),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    widths = (Inches(1.28), Inches(0.8), Inches(1.0), Inches(3.35))
    for i, width in enumerate(widths):
        table.columns[i].width = width
    headers = ["COLOR", "SWATCH", "VALUE", "PRIMARY USE"]
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, DEEP_FOREST)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        add_run(p, header, bold=True, size=8, color=WHITE)
        set_spacing(p, after=0)
    set_repeat_table_header(table.rows[0])
    for name, hex_value, value, use in rows:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell, top=100, bottom=100)
            set_cell_border(cell, color="D6DED5", size="4")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        add_run(cells[0].paragraphs[0], name, bold=True, size=8.8, color=CHARCOAL)
        set_cell_shading(cells[1], hex_value)
        add_run(cells[2].paragraphs[0], value, size=8.8, color=CHARCOAL)
        add_run(cells[3].paragraphs[0], use, size=8.5, color=CHARCOAL)
    doc.add_paragraph()


def add_two_column_cards(doc, cards):
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    for index in range(0, len(cards), 2):
        row = table.add_row()
        pair = cards[index:index + 2]
        for i in range(2):
            cell = row.cells[i]
            set_cell_margins(cell, top=150, start=160, bottom=150, end=160)
            set_cell_shading(cell, CREAM if index % 4 == 0 else PALE_LEAF)
            set_cell_border(cell, color="E1E5DB", size="4")
            if i < len(pair):
                title, body = pair[i]
                p = cell.paragraphs[0]
                add_run(p, title, bold=True, size=10.5, color=DEEP_FOREST, font="Georgia")
                set_spacing(p, after=4)
                p = cell.add_paragraph()
                add_run(p, body, size=9.2, color=CHARCOAL)
                set_spacing(p, after=0, line=1.05)
    doc.add_paragraph()


def add_page_number_footer(section):
    footer = section.footer
    for element in list(footer._element):
        footer._element.remove(element)
    p = footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "SEEDS OF LIFE GLOBAL  |  BRAND GUIDELINES  |  ", size=7.5, color=FOREST)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    p._p.append(fld_char1)
    p._p.append(instr_text)
    p._p.append(fld_char2)


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor.from_string(CHARCOAL)
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Georgia"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
        style.font.color.rgb = RGBColor.from_string(DEEP_FOREST)
    if "Brand Caption" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("Brand Caption", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Arial"
        style.font.size = Pt(8.3)
        style.font.color.rgb = RGBColor.from_string(MOSS)


def cover(doc):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    left, right = table.cell(0, 0), table.cell(0, 1)
    set_cell_shading(left, FOREST)
    set_cell_shading(right, CREAM)
    set_cell_margins(left, top=360, start=300, bottom=360, end=250)
    set_cell_margins(right, top=250, start=200, bottom=200, end=200)
    left.width, right.width = Inches(4.35), Inches(2.15)
    p = left.paragraphs[0]
    add_run(p, "SEEDS OF LIFE GLOBAL", bold=True, size=10, color=SEED_GOLD)
    set_spacing(p, after=11)
    p = left.add_paragraph()
    add_run(p, "Brand\nGuidelines", bold=True, size=29, color=WHITE, font="Georgia")
    set_spacing(p, after=14)
    p = left.add_paragraph()
    add_run(p, "A practical identity system for faith-rooted work with children, families, and communities.", size=11.2, color=WHITE)
    set_spacing(p, after=28, line=1.15)
    p = left.add_paragraph()
    add_run(p, "VERSION 1.0  |  SEPTEMBER 2026", bold=True, size=8.2, color=SEED_GOLD)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ICON_COLOR), width=Inches(1.68))
    cap = right.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(cap, "Where purpose takes root\nand lives bear fruit.", italic=True, size=10.4, color=DEEP_FOREST, font="Georgia")
    set_spacing(cap, before=15, after=0)
    doc.add_paragraph()
    add_rule(doc)
    add_para(doc, "This guide translates Seeds of Life Global's current visual identity and approved organizational messaging into consistent, respectful public communication.", size=9.6, color=MOSS, before=10, after=0)
    doc.add_page_break()


def contents_page(doc):
    add_kicker(doc, "Guide map")
    add_heading(doc, "Inside this guide")
    add_para(doc, "Use this document as the shared reference for public-facing work. It covers the essential choices that make Seeds of Life Global recognizable, respectful, and consistent.", size=10.7, after=12, line=1.14)
    entries = [
        ("01", "Brand foundation", "Mission, audience, and character"),
        ("02", "Logo system", "Approved marks, clear space, and misuse"),
        ("03", "Color", "Core palette and accessible use"),
        ("04", "Typography", "Type pairing and hierarchy"),
        ("05", "Voice and tone", "Writing principles and examples"),
        ("06", "Messaging and impact", "Message pillars and claim language"),
        ("07", "Photography and visual content", "Dignity, safeguarding, and captions"),
        ("08", "Partners and digital", "Partner presentation and accessibility"),
        ("09", "Application and governance", "Review checklist and updates"),
    ]
    table = doc.add_table(rows=0, cols=3)
    table.autofit = False
    for num, title, detail in entries:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell, top=105, bottom=105, start=120, end=120)
            set_cell_border(cell, color="D6DED5", size="4")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cells[0], FOREST)
        add_run(cells[0].paragraphs[0], num, bold=True, size=9, color=SEED_GOLD)
        add_run(cells[1].paragraphs[0], title, bold=True, size=10, color=DEEP_FOREST, font="Georgia")
        add_run(cells[2].paragraphs[0], detail, size=9, color=MOSS)
    doc.add_paragraph()
    add_callout(doc, "Working guide", "This is a living guide. Use the approved mission and visual system now, then revise the source file when new official assets, reporting periods, or campaign decisions are approved.", fill=PALE_LEAF)
    doc.add_page_break()


def main():
    for asset in (FULL_LOGO, ICON_COLOR, ICON_FOREST, ICON_WHITE):
        if not asset.exists():
            raise FileNotFoundError(f"Required brand asset missing: {asset}")
    doc = Document(str(REFERENCE))
    clear_body(doc)
    configure_styles(doc)
    for section in doc.sections:
        add_page_number_footer(section)

    cover(doc)
    contents_page(doc)

    # 1. Foundation
    add_kicker(doc, "01  /  Brand foundation")
    add_heading(doc, "Rooted in faith. Designed for real life.")
    add_para(doc, "Seeds of Life Global is a faith-rooted non-profit serving children, youth, families, and the people who care for them. The brand should feel as alive as a growing seed: hopeful, grounded, practical, and full of possibility.", size=11, after=8, line=1.16)
    add_callout(doc, "Mission", "At Seeds of Life Global, our mission is to nurture a sense of purpose within every child so that they may grow in faith, character, and confidence to make a difference in the world around them. Rooted in John 15:5, we believe lasting growth and impact come from abiding in Christ.", fill=CREAM)
    add_heading(doc, "Who we serve", level=2)
    add_two_column_cards(doc, [
        ("Children and youth", "Especially children in underserved communities who need safe, encouraging spaces to learn, belong, and grow."),
        ("Parents and families", "Caregivers who benefit from community, faith formation, practical support, and opportunity."),
        ("Churches, schools, and leaders", "Local leaders who equip children and youth through teaching, mentorship, and consistent care."),
        ("Partners and supporters", "Individuals and organizations who want their giving, expertise, and networks to create durable community impact."),
    ])
    add_heading(doc, "Brand character", level=2)
    for text in [
        "Rooted: faith is the source of the work, not a decorative afterthought.",
        "Hopeful: show what is possible without denying real needs.",
        "Welcoming: speak to children, families, and partners with equal dignity.",
        "Purposeful: connect every activity to growth, belonging, and agency.",
        "Credible: use clear, verifiable language and specific evidence.",
    ]:
        add_bullet(doc, text)
    doc.add_page_break()

    # 2. Logo
    add_kicker(doc, "02  /  Logo system")
    add_heading(doc, "The mark carries the story.")
    add_para(doc, "The logo combines a seed, growing leaves, a globe-like form, and a cross. Together, these elements communicate Christian faith, nurturing, global care, and a life that grows outward. Treat it as one complete symbol; the proportions and relationship of its elements are fixed.", size=10.8, after=10, line=1.14)
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    cell_a, cell_b = table.cell(0, 0), table.cell(0, 1)
    for cell in (cell_a, cell_b):
        set_cell_margins(cell, top=180, start=160, bottom=180, end=160)
        set_cell_border(cell, color="D6DED5", size="4")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(cell_a, CREAM)
    cell_a.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell_a.paragraphs[0].add_run().add_picture(str(FULL_LOGO), width=Inches(3.55))
    set_cell_shading(cell_b, FOREST)
    cell_b.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell_b.paragraphs[0].add_run().add_picture(str(ICON_WHITE), width=Inches(1.55))
    p = cell_b.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "Use the icon for\nsmall or square spaces.", size=8.8, color=WHITE)
    doc.add_paragraph()
    add_heading(doc, "Approved logo use", level=2)
    add_two_column_cards(doc, [
        ("Primary lockup", "Use the full-color wordmark on cream or white when the organization needs to be identified clearly."),
        ("Icon mark", "Use the icon alone for social avatars, favicons, stickers, and compact applications where the name appears nearby."),
        ("One-color forest", "Use on light backgrounds only when full color is impractical, such as one-color print or embroidery."),
        ("One-color white", "Use on a dark forest field or a photograph with consistently dark, uncluttered contrast behind it."),
    ])
    add_heading(doc, "Clear space and size", level=2)
    add_bullet(doc, "Keep clear space around every logo equal to at least one-quarter of the icon's height. Keep copy, edges, and other marks outside this area.")
    add_bullet(doc, "Recommended minimum: full lockup 200 px wide on screen or 1.5 in wide in print; icon 32 px wide on screen or 0.35 in wide in print.")
    add_bullet(doc, "Use the supplied PNG or future approved vector file. Do not recreate the wordmark by typing the name in another font.")
    add_heading(doc, "Never", level=2)
    add_bullet(doc, "Stretch, rotate, outline, recolor, crop, or rearrange the mark.")
    add_bullet(doc, "Place the full-color logo on a busy photograph or a low-contrast field.")
    add_bullet(doc, "Add shadows, gradients, effects, taglines, or partner logos inside the logo clear space.")
    doc.add_page_break()

    # 3. Colour
    add_kicker(doc, "03  /  Color")
    add_heading(doc, "A living, grounded palette.")
    add_para(doc, "Forest and seed gold are the visual anchors from the current logo. Deep forest, cream, moss, clay, and pale leaf support accessible, warm layouts without competing with the mark.", size=10.8, after=10, line=1.14)
    add_color_swatch_table(doc)
    add_callout(doc, "Accessibility", "For body copy, default to deep forest on cream or white. Check every final text-and-background pairing to WCAG AA contrast before publishing. Seed gold is an accent, not a small-text color on white.", fill=CREAM)
    add_heading(doc, "Suggested balance", level=2)
    add_bullet(doc, "Light ground first: cream or white should carry most page area.")
    add_bullet(doc, "Forest is for structure: headers, key panels, and brand moments.")
    add_bullet(doc, "Seed gold is for energy: a rule, detail, badge, or small highlight.")
    add_bullet(doc, "Use clay sparingly for a clear call to action or human-centered emphasis.")
    doc.add_page_break()

    # 4. Typography
    add_kicker(doc, "04  /  Typography")
    add_heading(doc, "Warm authority, plainspoken clarity.")
    add_para(doc, "Use a classic serif for meaningful, reflective headlines and a familiar sans-serif for clear everyday reading. This pairing makes faith-centered language feel human and readable instead of formal or distant.", size=10.8, after=10, line=1.14)
    add_two_column_cards(doc, [
        ("Display: Georgia", "Use bold Georgia for page titles, campaign headlines, pull quotes, and section headings. Sentence case is preferred."),
        ("Body: Arial / Helvetica", "Use regular or semibold Arial/Helvetica for paragraphs, buttons, captions, labels, forms, and navigation."),
        ("Wordmark", "The Seeds of Life Global name in the logo is artwork, not a type treatment. Always use an approved logo file."),
        ("Fallbacks", "If the preferred fonts are unavailable, use Times New Roman for display and a system sans-serif for body copy."),
    ])
    add_heading(doc, "Recommended digital scale", level=2)
    table = doc.add_table(rows=1, cols=4)
    for i, h in enumerate(["ROLE", "FAMILY / WEIGHT", "SIZE / LEADING", "USE"]):
        cell = table.cell(0, i); set_cell_shading(cell, DEEP_FOREST); set_cell_margins(cell)
        add_run(cell.paragraphs[0], h, bold=True, size=8, color=WHITE)
    set_repeat_table_header(table.rows[0])
    for role, family, scale, use in [
        ("Display 1", "Georgia Bold", "52-64 / 1.05", "Hero headlines"),
        ("Display 2", "Georgia Bold", "34-44 / 1.12", "Section headings"),
        ("Heading", "Arial Bold", "20-26 / 1.2", "Cards and subheads"),
        ("Body", "Arial Regular", "16-18 / 1.55", "Web reading"),
        ("Label", "Arial Bold", "12-14 / 1.2", "Navigation and metadata"),
    ]:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell, top=90, bottom=90); set_cell_border(cell, color="D6DED5", size="4")
        for cell, value in zip(cells, (role, family, scale, use)):
            add_run(cell.paragraphs[0], value, size=8.6, color=CHARCOAL, bold=(cell is cells[0]))
    doc.add_paragraph()
    add_callout(doc, "Typesetting", "Use sentence case, generous line height, and short paragraphs. Avoid blocks of centered text, all-caps body copy, and more than two type families in one asset.", fill=PALE_LEAF)
    doc.add_page_break()

    # 5. Voice
    add_kicker(doc, "05  /  Voice and tone")
    add_heading(doc, "Speak with hope and evidence.")
    add_para(doc, "Seeds of Life Global speaks from Christian conviction with humility. The voice invites people into a shared work; it does not speak down to communities or promise what the organization cannot verify.", size=10.8, after=10, line=1.14)
    add_two_column_cards(doc, [
        ("Warm, not vague", "Name people, places, and outcomes when they are approved. Replace generic inspiration with a concrete invitation."),
        ("Faith-rooted, not exclusive", "Let Scripture and faith inform the language while keeping invitations understandable and welcoming to every reader."),
        ("Dignifying, not paternalistic", "Center children's strengths, families' agency, and local leaders' work. Avoid language that portrays people as helpless."),
        ("Clear, not corporate", "Use short, active sentences. Prefer 'join a Bible study' over abstract phrases such as 'activate community transformation.'"),
    ])
    add_heading(doc, "Writing examples", level=2)
    table = doc.add_table(rows=1, cols=2)
    for i, h in enumerate(["PREFER", "AVOID"]):
        cell = table.cell(0, i); set_cell_shading(cell, FOREST if i == 0 else CLAY); set_cell_margins(cell)
        add_run(cell.paragraphs[0], h, bold=True, size=8.2, color=WHITE)
    set_repeat_table_header(table.rows[0])
    examples = [
        ("Children are growing in faith, character, and confidence through consistent community care.", "We save children from poverty."),
        ("Join us as we create spaces where purpose can take root.", "Help us change the world forever."),
        ("More than 1,000 children have participated in ministry and feeding programs around Sibonga, Cebu.", "We have impacted countless lives everywhere."),
        ("Your partnership helps local leaders keep showing up for children and families.", "Your gift guarantees success."),
    ]
    for good, avoid in examples:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell, top=110, bottom=110); set_cell_border(cell, color="D6DED5", size="4")
        add_run(cells[0].paragraphs[0], good, size=9.1, color=CHARCOAL)
        add_run(cells[1].paragraphs[0], avoid, size=9.1, color=CHARCOAL)
    doc.add_paragraph()
    add_heading(doc, "Calls to action", level=2)
    add_bullet(doc, "Good options: Grow with us. Explore our purpose. Join a Bible study. Partner in hope. Connect with Seeds of Life.")
    add_bullet(doc, "Pair a call to action with the next step: donate, volunteer, attend, learn more, or contact the team.")
    doc.add_page_break()

    # 6 Messaging and impact
    add_kicker(doc, "06  /  Messaging and impact")
    add_heading(doc, "One story. Four proof points.")
    add_para(doc, "Use this message hierarchy to keep different audiences connected to the same purpose. Adapt the emphasis, but do not change the heart of the story.", size=10.8, after=8, line=1.14)
    add_two_column_cards(doc, [
        ("Faith that roots growth", "Lasting growth begins in Christ. Scripture, belonging, and everyday encouragement help children discover purpose."),
        ("Children and families who flourish", "Programs create practical spaces for children, youth, and parents to learn, connect, and grow together."),
        ("Local leaders who multiply care", "Teachers, pastors, and youth leaders extend the work through training, community Bible studies, and faithful presence."),
        ("Places built for the future", "Mt. Moriah is being developed as a gathering and camping space designed to support children's ministry and community life."),
    ])
    add_heading(doc, "Current impact language", level=2)
    add_callout(doc, "Use with date context", "The figures below reflect organization-provided information for the current brand guide. Before a campaign, annual report, grant, or donor page, confirm the reporting period and source with leadership.", fill=CREAM)
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["APPROVED FIGURE", "CONTEXT", "SAFE PUBLIC WORDING"]):
        cell = table.cell(0, i); set_cell_shading(cell, DEEP_FOREST); set_cell_margins(cell)
        add_run(cell.paragraphs[0], h, bold=True, size=7.7, color=WHITE)
    set_repeat_table_header(table.rows[0])
    for figure, context, wording in [
        ("1,000+ children", "Children's ministry and feeding programs around Sibonga, Cebu.", "More than 1,000 children have participated in Seeds of Life Global programs around Sibonga, Cebu."),
        ("500+ leaders", "Teachers, pastors, and youth leaders trained in the Philippines and parts of Asia through Community Bible Study Children and Youth curriculum.", "More than 500 leaders have been trained to serve children and youth."),
        ("100+ communities", "Churches, schools, and communities reached across the Philippines and parts of Asia.", "Seeds of Life Global has reached more than 100 communities, churches, and schools."),
        ("Mt. Moriah", "Construction stage; one cabin is being built and can host gatherings and camping activities while work continues.", "Mt. Moriah is in construction and is beginning to welcome gatherings and camping activities."),
    ]:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell, top=105, bottom=105); set_cell_border(cell, color="D6DED5", size="4")
        add_run(cells[0].paragraphs[0], figure, bold=True, size=8.6, color=DEEP_FOREST)
        add_run(cells[1].paragraphs[0], context, size=8.25, color=CHARCOAL)
        add_run(cells[2].paragraphs[0], wording, size=8.25, color=CHARCOAL)
    doc.add_paragraph()
    doc.add_page_break()

    # 7 Image
    add_kicker(doc, "07  /  Photography and visual content")
    add_heading(doc, "Show participation, dignity, and place.")
    add_para(doc, "Images should make the work feel immediate and truthful. Use real moments of learning, play, prayer, leadership, family connection, and local context. The goal is not to manufacture emotion; it is to help people see community care in action.", size=10.8, after=9, line=1.14)
    add_two_column_cards(doc, [
        ("Choose", "Natural light, eye-level moments, activity in context, respectful portraits, and images that show children's agency and joy."),
        ("Avoid", "Humiliating or distressing images, staged hardship, crowded backgrounds behind the logo, and photos that identify children without permission."),
        ("Crop", "Use waist-up or head-and-shoulders leadership portraits for balanced team cards. Leave simple negative space for copy only when it does not hide the story."),
        ("Caption", "Include what is happening, a general location, and a date or reporting period when practical. Example: 'Children's Bible Study, Sibonga, Cebu - 2026.'"),
    ])
    add_heading(doc, "Consent and safeguarding", level=2)
    add_bullet(doc, "Obtain and record consent appropriate to the person and setting before public use. For children, follow the organization's safeguarding and parent/guardian-permission process.")
    add_bullet(doc, "Do not include full names, medical information, school details, or locations that could put a child at risk without explicit approval.")
    add_bullet(doc, "Store the original photo, permission record, caption, location, date, and credit together so future use stays accurate.")
    add_callout(doc, "Logo on photography", "Use the white icon only on a consistently dark, simple image area. If the image needs a large logo panel to be readable, use a separate forest panel rather than forcing the mark over the photo.", fill=PALE_LEAF)
    doc.add_page_break()

    # 8 Partners & digital
    add_kicker(doc, "08  /  Partners and digital")
    add_heading(doc, "Make every relationship clear.")
    add_para(doc, "Partner names and logos strengthen credibility when they are accurate, current, and authorized. Treat each partner's identity with the same care given to Seeds of Life Global's own mark.", size=10.8, after=9, line=1.14)
    add_heading(doc, "Partner presentation", level=2)
    add_bullet(doc, "Use only approved partner names and logo files. Confirm written permission, correct spelling, and current relationship before release.")
    add_bullet(doc, "Give partner logos equal visual height and generous clear space. Do not make a partner look like an endorser unless that is explicitly true.")
    add_bullet(doc, "Current publicly approved names: Community Bible Study; New Life Assembly of God Church - Sibonga; Skyrise Realty and Development Corporation; The Maclellan Foundation, Inc.; and Brainster Academy.")
    add_heading(doc, "Digital essentials", level=2)
    add_two_column_cards(doc, [
        ("Website", "Use the full lockup in the header and footer where space allows. Keep navigation simple and calls to action specific."),
        ("Social profiles", "Use the icon mark for avatars. Use the full lockup in headers, profile graphics, and pinned brand materials."),
        ("Accessible content", "Write descriptive alt text. Example: 'Seeds of Life Global seed-and-cross logo in gold and forest green.' Add captions to video and avoid text baked into images when HTML text will work."),
        ("Email and documents", "Use a small full-color lockup, deep forest headings, and short blocks of content. Keep the signature, links, and phone numbers selectable as text."),
    ])
    add_heading(doc, "Recommended alt text", level=2)
    add_callout(doc, "Logo", "Seeds of Life Global logo: a gold seed and leaves above a globe form with a green cross, followed by the name Seeds of Life Global.", fill=CREAM)
    add_callout(doc, "Program photo", "Children participating in a Bible study activity in Sibonga, Cebu.", fill=PALE_LEAF)
    doc.add_page_break()

    # 9 Application and governance
    add_kicker(doc, "09  /  Application and governance")
    add_heading(doc, "A simple check before anything goes live.")
    add_para(doc, "Consistency matters because trust is built through small repeated signals. Use this checklist for web pages, social posts, print, fundraising materials, presentations, and partner announcements.", size=10.8, after=9, line=1.14)
    add_heading(doc, "Brand review checklist", level=2)
    checklist = [
        "The approved logo file, color version, clear space, and minimum size are used.",
        "Color pairings and body text meet accessible contrast standards.",
        "Headlines are in the display serif; reading copy is in the approved sans-serif.",
        "Language is warm, clear, truthful, and respectful of children, families, and local leaders.",
        "Impact figures include the right context and have been verified for the reporting period.",
        "Photos are dignified, accurately captioned, and cleared for public use.",
        "Partner names, logos, and descriptions are current and approved.",
        "Links, phone numbers, addresses, donation instructions, and contact details have been tested.",
    ]
    for item in checklist:
        p = doc.add_paragraph()
        add_run(p, "□  ", bold=True, size=12, color=FOREST)
        add_run(p, item, size=10.1, color=CHARCOAL)
        set_spacing(p, after=5, line=1.08)
    add_heading(doc, "Asset library", level=2)
    add_para(doc, "Use the approved files in the organization's brand folder. Request a vector logo (SVG, EPS, or AI) from the original designer before large-format print, signage, or embroidery production. The current PNG files are appropriate for common digital use when displayed at sensible sizes.", size=10.3, after=8, line=1.12)
    add_heading(doc, "Keeping this guide current", level=2)
    add_bullet(doc, "Review impact figures, program status, partners, contact details, and photo-consent practices at least annually and before major campaigns.")
    add_bullet(doc, "Update this guide when an approved vector logo, new campaign identity, or revised organizational messaging is adopted.")
    add_callout(doc, "Approval", "Brand decisions, public impact claims, child-related imagery, and partner use should be approved by Seeds of Life Global leadership before publication.", fill=FOREST, label_color=SEED_GOLD)
    # Re-color body run in last callout for dark field.
    last_cell = doc.tables[-1].cell(0, 0)
    for run in last_cell.paragraphs[0].runs:
        if run.text.strip() and not run.bold:
            run.font.color.rgb = RGBColor.from_string(WHITE)

    doc.core_properties.title = "Seeds of Life Global Brand Guidelines"
    doc.core_properties.subject = "Visual identity, voice, and communications standards"
    doc.core_properties.author = "Seeds of Life Global"
    doc.core_properties.comments = "Prepared from approved organization-provided brand information."
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
